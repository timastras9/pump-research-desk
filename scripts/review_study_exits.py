#!/usr/bin/env python3
"""Add a 'Your study' sheet to the exit-policy workbook: for each token recorded in a study, the most profitable
sale point (hindsight high, filled after the latency), what the live paper trader did, and what the model would do.

  python scripts/review_study_exits.py --db artifacts/corpus/launches.db --model-dir artifacts/corpus/exit-policy-v2 \
      --study-tokens study_tokens.json --picks jonesy,shibu,pomkori,waste,life,yap
"""
import argparse, json, os, pickle, sys, time
import numpy as np
sys.path.insert(0, os.path.dirname(__file__))
import train_exit_policy as tp

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--db', required=True); ap.add_argument('--model-dir', required=True); ap.add_argument('--study-tokens', required=True, help='wrangler d1 --json export of study_tokens rows (id,data)')
    ap.add_argument('--picks', default=''); ap.add_argument('--latency', type=int, default=2)
    a = ap.parse_args(); tp.LATENCY_S = a.latency
    res = json.load(open(os.path.join(a.model_dir, 'results.json'))); th = res['threshold_chosen_on_validation']
    model = pickle.load(open(os.path.join(a.model_dir, 'model.pkl'), 'rb'))
    study = [json.loads(r['data']) for r in json.load(open(a.study_tokens))[0]['results']]
    corpus = {t['mint']: t for t in tp.load(a.db, 1)}
    picks = [p.strip().lower() for p in a.picks.split(',') if p.strip()]
    policy = lambda t, e, s: model.predict_proba(np.array([tp.features(t, e, s)], dtype=np.float32))[0, 1] >= th or (t['p'][s] / t['p'][e] - 1) <= -0.25
    rows = []
    for st in study:
        t = corpus.get(st['mint']); paper = st.get('paper') or {}
        pick = any(p in (st.get('name') or '').lower() for p in picks)
        base = [('★ ' if pick else '') + (st.get('name') or ''), st['mint'], st.get('status'), paper.get('status'), paper.get('exitReason') or paper.get('skipReason'), paper.get('pnlPct')]
        if not t: rows.append((pick, base + ['not in corpus yet (candles fetched ~13 min after launch)'] + [None] * 12)); continue
        t['entry'] = tp.entry_of(t); e = t['entry']
        if e is None: rows.append((pick, base + ['no entry: no trade within 30 s of first sight, or already +30%'] + [None] * 12)); continue
        p = t['p']; end = min(720, e + tp.HOLD_S); k = e + 1 + int(np.argmax(p[e + 1:end + 1])); kf = min(k + tp.LATENCY_S, end)
        md, mf, mp = tp.simulate(t, policy, detail=True)
        rows.append((pick, base + ['', e, p[e], k, round((p[k] / p[e] - 1) * 100, 1), kf, round(tp.net(p[e], p[kf]), 1), round(0.02 * tp.net(p[e], p[kf]), 2), md, mf, round(mp, 1), round(0.02 * mp, 2), round((p[mf] / p[k] - 1) * 100, 1)]))
    rows.sort(key=lambda r: (not r[0], -(r[1][12] if isinstance(r[1][12], (int, float)) else -1e9)))
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    path = os.path.join(a.model_dir, 'exit_policy_review.xlsx'); wb = load_workbook(path)
    if 'Your study' in wb.sheetnames: del wb['Your study']
    ws = wb.create_sheet('Your study', 1)
    header = ['Token (★ = your pick)', 'Mint', 'Recording', 'Live paper trade', 'Live exit / skip reason', 'Live paper P&L %', 'Note', 'Entry sec after launch', 'Entry price',
              'Most profitable sale: HIGH sec', 'HIGH % vs entry', f'Fill sec (+{tp.LATENCY_S}s)', 'Best realistic P&L %', 'Best realistic $ on $2', 'Model decided sec', 'Model fill sec', 'Model P&L %', 'Model $ on $2', 'Model fill vs HIGH %']
    ws.append(header)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7')
    for _, r in rows: ws.append(r)
    ws.freeze_panes = 'A2'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(10, min(36, max(len(str(c.value or '')) for c in col) + 2))
    wb.save(path)
    for pick, r in rows:
        if pick: print(json.dumps(dict(zip(header, r)), default=str))

if __name__ == '__main__':
    main()
